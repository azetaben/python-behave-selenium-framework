"""Utility for loading test data from external sources (JSON, CSV, Excel)."""

import csv
import json
from pathlib import Path
from typing import List, Dict, Any

try:
    import openpyxl
    HAS_OPENPYXL = True
except ImportError:
    HAS_OPENPYXL = False

from models import ExternalLoginDataRow
from constants.app_constants import AppConstants


class TestDataLoader:
    """Load and parse test data from external sources."""

    @staticmethod
    def load_login_json(json_file_path: str) -> List[ExternalLoginDataRow]:
        """
        Load login test data from JSON file.

        Expected JSON format:
        [
            {
                "testCaseId": "TC_001",
                "username": "user",
                "password": "pass",
                "expectedResult": "SUCCESS",
                "expectedMessage": ""
            },
            ...
        ]

        Args:
            json_file_path: Path to JSON file

        Returns:
            List of ExternalLoginDataRow objects
        """
        path = Path(json_file_path)
        if not path.exists():
            raise FileNotFoundError(f"JSON file not found: {json_file_path}")

        with open(path, 'r', encoding='utf-8') as f:
            data = json.load(f)

        rows = []
        for item in data:
            row = ExternalLoginDataRow(
                test_case_id=item.get('testCaseId', ''),
                username=item.get('username', ''),
                password=item.get('password', ''),
                expected_result=item.get('expectedResult', ''),
                expected_message=item.get('expectedMessage', '')
            )
            rows.append(row)

        return rows

    @staticmethod
    def load_login_csv(csv_file_path: str) -> List[ExternalLoginDataRow]:
        """
        Load login test data from CSV file.

        Expected CSV format (header required):
        test_case_id,username,password,expected_result,expected_message
        TC_001,user,pass,SUCCESS,
        TC_002,invalid,wrong,FAILURE,Invalid credentials

        Args:
            csv_file_path: Path to CSV file

        Returns:
            List of ExternalLoginDataRow objects
        """
        path = Path(csv_file_path)
        if not path.exists():
            raise FileNotFoundError(f"CSV file not found: {csv_file_path}")

        rows = []
        with open(path, 'r', encoding='utf-8') as f:
            reader = csv.DictReader(f)
            for row_dict in reader:
                row = ExternalLoginDataRow(
                    test_case_id=row_dict.get('test_case_id', '').strip(),
                    username=row_dict.get('username', '').strip(),
                    password=row_dict.get('password', '').strip(),
                    expected_result=row_dict.get('expected_result', '').strip(),
                    expected_message=row_dict.get('expected_message', '').strip()
                )
                rows.append(row)

        return rows

    @staticmethod
    def load_login_excel(excel_file_path: str, sheet_name: str = 'LoginTestData') -> List[ExternalLoginDataRow]:
        """
        Load login test data from Excel file.

        Expected columns:
        test_case_id, username, password, expected_result, expected_message

        Args:
            excel_file_path: Path to Excel file
            sheet_name: Name of the sheet containing test data (default: 'LoginTestData')

        Returns:
            List of ExternalLoginDataRow objects

        Raises:
            ImportError: If openpyxl is not installed
            FileNotFoundError: If file doesn't exist
        """
        if not HAS_OPENPYXL:
            raise ImportError(
                "openpyxl is required to load Excel files. "
                "Install with: pip install openpyxl"
            )

        path = Path(excel_file_path)
        if not path.exists():
            raise FileNotFoundError(f"Excel file not found: {excel_file_path}")

        workbook = openpyxl.load_workbook(path)

        if sheet_name not in workbook.sheetnames:
            raise ValueError(
                f"Sheet '{sheet_name}' not found. Available sheets: {workbook.sheetnames}"
            )

        worksheet = workbook[sheet_name]
        rows = []

        # Get header row
        headers = []
        for cell in worksheet[1]:
            headers.append(cell.value)

        # Get data rows
        for row_idx, row in enumerate(worksheet.iter_rows(min_row=2, values_only=True), start=2):
            if all(cell is None for cell in row):
                continue  # Skip empty rows

            row_dict = dict(zip(headers, row))

            row_obj = ExternalLoginDataRow(
                test_case_id=str(row_dict.get('test_case_id', '')).strip(),
                username=str(row_dict.get('username', '')).strip(),
                password=str(row_dict.get('password', '')).strip(),
                expected_result=str(row_dict.get('expected_result', '')).strip(),
                expected_message=str(row_dict.get('expected_message', '')).strip()
            )
            rows.append(row_obj)

        workbook.close()
        return rows

    @staticmethod
    def load_login_data(file_path: str, file_format: str = None, **kwargs) -> List[ExternalLoginDataRow]:
        """
        Auto-detect and load test data from file.

        Args:
            file_path: Path to test data file
            file_format: Format of file ('json', 'csv', 'xlsx'). If None, auto-detected from extension.
            **kwargs: Additional arguments (e.g., sheet_name for Excel)

        Returns:
            List of ExternalLoginDataRow objects
        """
        if file_format is None:
            # Auto-detect format from file extension
            path = Path(file_path)
            ext = path.suffix.lower()
            if ext == '.json':
                file_format = 'json'
            elif ext == '.csv':
                file_format = 'csv'
            elif ext in ['.xlsx', '.xls']:
                file_format = 'xlsx'
            else:
                raise ValueError(f"Unknown file format: {ext}")

        if file_format.lower() == 'json':
            return TestDataLoader.load_login_json(file_path)
        elif file_format.lower() == 'csv':
            return TestDataLoader.load_login_csv(file_path)
        elif file_format.lower() == 'xlsx':
            return TestDataLoader.load_login_excel(file_path, **kwargs)
        else:
            raise ValueError(f"Unsupported file format: {file_format}")

    @staticmethod
    def get_test_data_path(filename: str) -> str:
        """
        Get path to test data file in testData directory.

        Args:
            filename: Name of test data file (e.g., 'loginData.json')

        Returns:
            Absolute path to test data file
        """
        base_path = Path(__file__).parent.parent / 'testData'
        return str(base_path / filename)

    @staticmethod
    def create_sample_csv(output_path: str) -> None:
        """
        Create a sample CSV file with login test data.

        Args:
            output_path: Path where to create the sample CSV
        """
        sample_data = [
            {
                'test_case_id': 'CSV_TC_001',
                'username': AppConstants.Users.STANDARD,
                'password': AppConstants.Users.PASSWORD,
                'expected_result': 'SUCCESS',
                'expected_message': ''
            },
            {
                'test_case_id': 'CSV_TC_002',
                'username': AppConstants.Users.LOCKED,
                'password': AppConstants.Users.PASSWORD,
                'expected_result': 'FAILURE',
                'expected_message': AppConstants.Errors.LOCKED_OUT_USER,
            },
            {
                'test_case_id': 'CSV_TC_003',
                'username': 'invalid_user',
                'password': 'wrong_password',
                'expected_result': 'FAILURE',
                'expected_message': 'Epic sadface: Username and password do not match any user in this service'
            },
            {
                'test_case_id': 'CSV_TC_004',
                'username': '',
                'password': AppConstants.Users.PASSWORD,
                'expected_result': 'FAILURE',
                'expected_message': AppConstants.Errors.USERNAME_REQUIRED,
            },
        ]

        path = Path(output_path)
        path.parent.mkdir(parents=True, exist_ok=True)

        with open(path, 'w', newline='', encoding='utf-8') as f:
            writer = csv.DictWriter(f, fieldnames=[
                'test_case_id', 'username', 'password', 'expected_result', 'expected_message'
            ])
            writer.writeheader()
            writer.writerows(sample_data)

    @staticmethod
    def create_sample_json(output_path: str) -> None:
        """
        Create a sample JSON file with login test data.

        Args:
            output_path: Path where to create the sample JSON
        """
        sample_data = [
            {
                'testCaseId': 'JSON_SAMPLE_001',
                'username': AppConstants.Users.STANDARD,
                'password': AppConstants.Users.PASSWORD,
                'expectedResult': 'SUCCESS',
                'expectedMessage': ''
            },
            {
                'testCaseId': 'JSON_SAMPLE_002',
                'username': AppConstants.Users.PROBLEM,
                'password': AppConstants.Users.PASSWORD,
                'expectedResult': 'SUCCESS',
                'expectedMessage': ''
            },
            {
                'testCaseId': 'JSON_SAMPLE_003',
                'username': 'invalid_user',
                'password': 'wrong_password',
                'expectedResult': 'FAILURE',
                'expectedMessage': AppConstants.Errors.WRONG_CREDENTIALS,
            },
        ]

        path = Path(output_path)
        path.parent.mkdir(parents=True, exist_ok=True)

        with open(path, 'w', encoding='utf-8') as f:
            json.dump(sample_data, f, indent=2, ensure_ascii=False)











